"""Builds orders_cleanup_demo.xlsx: messy order export -> formula-driven clean
sheet -> summary with a reconciliation check. Every formula is plain
Excel/Google Sheets (TRIM, PROPER, IFERROR, INDEX/MATCH, COUNTIFS, SUMIFS),
so the same workbook opens and recalculates in either.

Run: python build_workbook.py  (needs openpyxl)
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAW = [  # date, customer (as exported), region, product, qty, unit price
    (dt.date(2026, 9, 1), "  acme hardware", "North", "Slicer Pro", 12, 14.50),
    (dt.date(2026, 9, 1), "Blue Oak Kitchen ", "South", "Peeler", 30, 3.25),
    (dt.date(2026, 9, 3), "ACME HARDWARE", "North", "Peeler", 20, 3.25),
    (dt.date(2026, 9, 4), "corner cook shop", "West", "Slicer Pro", 6, 14.50),
    (dt.date(2026, 9, 4), "corner cook shop", "West", "Slicer Pro", 6, 14.50),  # duplicate
    (dt.date(2026, 9, 8), "Blue Oak kitchen", "South", "Grater Mini", 18, 5.75),
    (dt.date(2026, 9, 9), "Harbor Home", "East", "Slicer Pro", 10, 14.50),
    (dt.date(2026, 9, 12), "harbor home", "East", "Herb Scissors", 15, 7.00),  # not in price list
    (dt.date(2026, 9, 15), "Acme Hardware", "North", "Grater Mini", 24, 5.75),
    (dt.date(2026, 9, 18), "Pine & Co", "West", "Peeler", 40, 3.25),
    (dt.date(2026, 9, 22), "pine & co", "West", "Slicer Pro", 8, 14.50),
    (dt.date(2026, 9, 25), "Corner Cook Shop", "West", "Grater Mini", 12, 5.75),
]
PRODUCTS = [("Slicer Pro", "Cutting"), ("Peeler", "Prep"), ("Grater Mini", "Prep")]
REGIONS = ["North", "South", "East", "West"]
CATEGORIES = ["Cutting", "Prep"]
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDE7E1")


def head(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD


def build(path="orders_cleanup_demo.xlsx"):
    wb = Workbook()
    raw = wb.active
    raw.title = "Raw"
    head(raw, ["Date", "Customer", "Region", "Product", "Qty", "Unit price"])
    for r in RAW:
        raw.append(list(r))
    n = len(RAW) + 1  # last data row

    lk = wb.create_sheet("Lookup")
    head(lk, ["Product", "Category"])
    for p in PRODUCTS:
        lk.append(list(p))
    m = len(PRODUCTS) + 1

    cl = wb.create_sheet("Clean")
    head(cl, ["Date", "Customer", "Region", "Product", "Category", "Qty", "Amount",
              "Duplicate", "Issue"])
    for i in range(2, n + 1):
        cl.append([
            f"=Raw!A{i}",
            f"=PROPER(TRIM(Raw!B{i}))",
            f"=Raw!C{i}",
            f"=Raw!D{i}",
            f'=IFERROR(INDEX(Lookup!$B$2:$B${m},MATCH(D{i},Lookup!$A$2:$A${m},0)),"UNMAPPED")',
            f"=Raw!E{i}",
            f"=Raw!E{i}*Raw!F{i}",
            # second and later copies of the same date/customer/product/qty are flagged
            f'=IF(COUNTIFS($A$2:A{i},A{i},$B$2:B{i},B{i},$D$2:D{i},D{i},$F$2:F{i},F{i})>1,"DUPLICATE","")',
            f'=IF(H{i}<>"","duplicate row",IF(E{i}="UNMAPPED","product not in Lookup",""))',
        ])
        cl[f"A{i}"].number_format = "yyyy-mm-dd"
        cl[f"G{i}"].number_format = "#,##0.00"

    red = PatternFill("solid", fgColor="F8D7D3")
    cl.conditional_formatting.add(f"A2:I{n}", FormulaRule(formula=["$I2<>\"\""], fill=red))
    dv = DataValidation(type="list", formula1='"' + ",".join(REGIONS) + '"', allow_blank=False)
    raw.add_data_validation(dv)
    dv.add(f"C2:C{n + 200}")

    sm = wb.create_sheet("Summary")
    head(sm, ["Region"] + CATEGORIES + ["Total"])
    # duplicates and unmapped rows are excluded from revenue, never silently counted
    for r, reg in enumerate(REGIONS, start=2):
        row = [reg]
        for c in CATEGORIES:
            row.append(f'=SUMIFS(Clean!$G$2:$G${n},Clean!$C$2:$C${n},$A{r},'
                       f'Clean!$E$2:$E${n},"{c}",Clean!$H$2:$H${n},"")')
        row.append(f"=SUM(B{r}:{chr(65 + len(CATEGORIES))}{r})")
        sm.append(row)
    t = len(REGIONS) + 2
    tot = chr(66 + len(CATEGORIES))
    sm.append(["Total"] + [f"=SUM({chr(66 + k)}2:{chr(66 + k)}{t - 1})" for k in range(len(CATEGORIES))]
              + [f"=SUM({tot}2:{tot}{t - 1})"])
    for c in sm[t]:
        c.font = BOLD
    sm.append([])
    sm.append(["Rows in export", f"=COUNTA(Raw!A2:A{n})"])
    sm.append(["Duplicate rows excluded", f'=COUNTIF(Clean!H2:H{n},"DUPLICATE")'])
    sm.append(["Unmapped rows excluded", f'=COUNTIFS(Clean!E2:E{n},"UNMAPPED",Clean!H2:H{n},"")'])
    sm.append(["Duplicate amount", f'=SUMIFS(Clean!G2:G{n},Clean!H2:H{n},"DUPLICATE")'])
    sm.append(["Unmapped amount", f'=SUMIFS(Clean!G2:G{n},Clean!E2:E{n},"UNMAPPED",Clean!H2:H{n},"")'])
    # independent check: summary total + both exclusions must equal the raw export
    sm.append(["Reconciliation", f'=IF(ABS(SUMPRODUCT(Raw!E2:E{n},Raw!F2:F{n})-{tot}{t}'
                                 f'-B{t + 5}-B{t + 6})<0.005,"OK","CHECK")'])
    for row in sm.iter_rows(min_row=2, max_row=t, min_col=2, max_col=1 + len(CATEGORIES) + 1):
        for c in row:
            c.number_format = "#,##0.00"
    for r in (t + 5, t + 6):
        sm[f"B{r}"].number_format = "#,##0.00"
    for ws, widths in ((raw, [12, 22, 10, 16, 6, 11]), (cl, [12, 20, 10, 16, 11, 6, 11, 12, 22]),
                       (sm, [24, 12, 12, 12]), (lk, [16, 12])):
        for k, w in enumerate(widths):
            ws.column_dimensions[chr(65 + k)].width = w
    wb.move_sheet("Summary", offset=-3)
    wb.save(path)
    return path


if __name__ == "__main__":
    print(build())
